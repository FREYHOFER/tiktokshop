#!/usr/bin/env python3
"""Create TikTok Shop products from the prepared Libri upload workbook.

This uses the TikTok Shop Open API path instead of Seller Center bulk upload:
1. download public Libri/media image URLs,
2. upload each image to TikTok Shop image storage,
3. create the product as LISTING or AS_DRAFT.

The script logs every attempted product and skips existing Seller SKUs by
default, so reruns are idempotent enough for operations work.
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import html
import json
import mimetypes
import re
import time
import urllib.error
import urllib.parse
import urllib.request
import uuid
from pathlib import Path
from typing import Any

import openpyxl

from fetch_libri_product_pages import load_env_file
from tiktok_order_automation import TikTokApiError, TikTokShopClient, clean, compact_json


DEFAULT_WORKBOOK = Path("outputs/catalog_expansion/20260617_1910/upload_pack_final/tiktok_upload_green.xlsx")
DEFAULT_OUTPUT_DIR = Path("outputs/catalog_expansion/20260617_1910/api_upload")
DEFAULT_CATEGORY_ID = "987016"  # Buecher, Zeitschriften und Audio > Literatur und Kunst > Roman
DEFAULT_WAREHOUSE_ID = "7630853807695791895"
DEFAULT_RESPONSIBLE_PERSON_ID = "69e8ebf9401201b6c3954a0f"
DEFAULT_WARNING_ATTRIBUTE_ID = "102277"
DEFAULT_WARNING_NO_VALUE_ID = "1000059"
DEFAULT_CURRENCY = "EUR"


class TikTokProductClient(TikTokShopClient):
    def upload_product_image(self, image_bytes: bytes, filename: str, use_case: str = "MAIN_IMAGE") -> dict[str, Any]:
        path = "/product/202309/images/upload"
        body, content_type = multipart_body(
            fields={"use_case": use_case},
            file_field="data",
            filename=filename,
            file_bytes=image_bytes,
            content_type=mimetypes.guess_type(filename)[0] or "image/jpeg",
        )

        query: dict[str, str] = {
            "app_key": self.app_key,
            "timestamp": str(int(time.time())),
        }
        if self.include_version:
            query["version"] = self.version
        if self.access_token_in_query:
            query["access_token"] = self.access_token
        query["sign"] = self.sign(path, query, body_text="", include_body=False)

        request = urllib.request.Request(
            self.base_url + path + "?" + urllib.parse.urlencode(query),
            data=body,
            method="POST",
            headers={
                "Content-Type": content_type,
                "x-tts-access-token": self.access_token,
                "User-Agent": "TikTokShop-Libri-API-Upload/1.0",
            },
        )
        try:
            with urllib.request.urlopen(request, timeout=90) as response:
                raw = response.read().decode("utf-8", errors="replace")
        except urllib.error.HTTPError as exc:
            detail = exc.read().decode("utf-8", errors="replace")
            raise TikTokApiError(f"TikTok image upload HTTP {exc.code}: {detail[:1000]}") from exc
        except urllib.error.URLError as exc:
            raise TikTokApiError(f"TikTok image upload connection failed: {exc}") from exc

        result = json.loads(raw)
        code = clean(result.get("code"))
        if code and code not in {"0", "OK", "SUCCESS"}:
            raise TikTokApiError(f"TikTok image upload error {code}: {clean(result.get('message')) or result}")
        return result.get("data") or {}


def multipart_body(
    fields: dict[str, str],
    file_field: str,
    filename: str,
    file_bytes: bytes,
    content_type: str,
) -> tuple[bytes, str]:
    boundary = f"----codex-tiktok-{uuid.uuid4().hex}"
    chunks: list[bytes] = []
    for key, value in fields.items():
        chunks.append(f"--{boundary}\r\n".encode("utf-8"))
        chunks.append(f'Content-Disposition: form-data; name="{key}"\r\n\r\n'.encode("utf-8"))
        chunks.append(str(value).encode("utf-8"))
        chunks.append(b"\r\n")
    chunks.append(f"--{boundary}\r\n".encode("utf-8"))
    chunks.append(
        (
            f'Content-Disposition: form-data; name="{file_field}"; filename="{filename}"\r\n'
            f"Content-Type: {content_type}\r\n\r\n"
        ).encode("utf-8")
    )
    chunks.append(file_bytes)
    chunks.append(b"\r\n")
    chunks.append(f"--{boundary}--\r\n".encode("utf-8"))
    return b"".join(chunks), f"multipart/form-data; boundary={boundary}"


def download_url(url: str) -> bytes:
    request = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(request, timeout=60) as response:
        return response.read()


def detect_headers(sheet) -> dict[str, int]:
    headers = {
        str(sheet.cell(row=1, column=col).value): col
        for col in range(1, sheet.max_column + 1)
        if sheet.cell(row=1, column=col).value
    }
    required = {"product_name", "product_description", "main_image", "gtin_code", "price", "quantity", "seller_sku"}
    missing = required - set(headers)
    if missing:
        raise ValueError(f"Workbook is missing required TikTok columns: {sorted(missing)}")
    return headers


def cell(sheet, headers: dict[str, int], row_idx: int, key: str) -> str:
    col = headers.get(key)
    if not col:
        return ""
    return clean(sheet.cell(row=row_idx, column=col).value)


def load_upload_rows(path: Path, limit: int = 0, start_row: int = 7) -> list[dict[str, Any]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["Template"] if "Template" in workbook.sheetnames else workbook.active
    headers = detect_headers(sheet)
    rows: list[dict[str, Any]] = []
    for row_idx in range(start_row, min(sheet.max_row, 5000) + 1):
        title = cell(sheet, headers, row_idx, "product_name")
        if not title:
            continue
        image_urls = [cell(sheet, headers, row_idx, "main_image")]
        for image_idx in range(2, 10):
            value = cell(sheet, headers, row_idx, f"image_{image_idx}")
            if value:
                image_urls.append(value)
        rows.append(
            {
                "row": row_idx,
                "title": title,
                "description": cell(sheet, headers, row_idx, "product_description"),
                "ean": re.sub(r"\D", "", cell(sheet, headers, row_idx, "gtin_code")),
                "price": cell(sheet, headers, row_idx, "price"),
                "quantity": cell(sheet, headers, row_idx, "quantity"),
                "seller_sku": cell(sheet, headers, row_idx, "seller_sku"),
                "parcel_weight_g": cell(sheet, headers, row_idx, "parcel_weight"),
                "parcel_length_cm": cell(sheet, headers, row_idx, "parcel_length"),
                "parcel_width_cm": cell(sheet, headers, row_idx, "parcel_width"),
                "parcel_height_cm": cell(sheet, headers, row_idx, "parcel_height"),
                "image_urls": list(dict.fromkeys(url for url in image_urls if url)),
            }
        )
        if limit and len(rows) >= limit:
            break
    return rows


def decimal_string(value: str) -> str:
    text = clean(value).replace(",", ".")
    number = float(text)
    return f"{number:.2f}".rstrip("0").rstrip(".")


def int_string(value: str, default: int) -> int:
    text = clean(value).replace(",", ".")
    if not text:
        return default
    return max(0, int(float(text)))


def description_to_html(value: str) -> str:
    lines = [clean(line) for line in str(value or "").splitlines()]
    paragraphs = [line for line in lines if line]
    if not paragraphs:
        return "<p>Produktdetails siehe Artikeldaten.</p>"
    return "".join(f"<p>{html.escape(line)}</p>" for line in paragraphs)


def build_payload(row: dict[str, Any], image_uris: list[str], args: argparse.Namespace) -> dict[str, Any]:
    weight_kg = max(0.001, int_string(row["parcel_weight_g"], 500) / 1000)
    quantity = int_string(row["quantity"], 1)
    payload = {
        "save_mode": args.save_mode,
        "description": description_to_html(row["description"]),
        "category_id": args.category_id,
        "main_images": [{"uri": uri} for uri in image_uris],
        "skus": [
            {
                "inventory": [{"warehouse_id": args.warehouse_id, "quantity": quantity}],
                "seller_sku": row["seller_sku"],
                "price": {"amount": decimal_string(row["price"]), "currency": args.currency},
                "identifier_code": {"code": row["ean"], "type": "ISBN"},
            }
        ],
        "title": row["title"],
        "is_cod_allowed": False,
        "package_dimensions": {
            "length": str(int_string(row["parcel_length_cm"], 25)),
            "width": str(int_string(row["parcel_width_cm"], 18)),
            "height": str(int_string(row["parcel_height_cm"], 4)),
            "unit": "CENTIMETER",
        },
        "product_attributes": [
            {
                "id": DEFAULT_WARNING_ATTRIBUTE_ID,
                "values": [{"id": DEFAULT_WARNING_NO_VALUE_ID}],
            }
        ],
        "package_weight": {"value": f"{weight_kg:.3f}".rstrip("0").rstrip("."), "unit": "KILOGRAM"},
        "responsible_person_ids": [args.responsible_person_id],
        "manufacturer_ids": [],
        "listing_platforms": ["TIKTOK_SHOP"],
        "shipping_insurance_requirement": "NOT_SUPPORTED",
        "minimum_order_quantity": 1,
        "is_pre_owned": False,
        "category_version": "v2",
        "idempotency_key": f"libri-{row['ean']}-v1",
    }
    return payload


def existing_sku(client: TikTokProductClient, seller_sku: str) -> str:
    response = client.request(
        "POST",
        "/product/202309/products/search",
        params={"shop_cipher": client.ensure_shop_cipher(), "page_size": 10},
        body={"seller_skus": [seller_sku]},
    )
    products = (response.get("data") or {}).get("products") or []
    if not products:
        return ""
    return clean(products[0].get("id") or products[0].get("product_id"))


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = [
        "timestamp_utc",
        "row",
        "ean",
        "seller_sku",
        "title",
        "status",
        "product_id",
        "sku_id",
        "image_count",
        "message",
    ]
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=fieldnames, delimiter=";")
        writer.writeheader()
        writer.writerows(rows)


def process(args: argparse.Namespace) -> int:
    env_path = Path(args.env)
    client = TikTokProductClient(load_env_file(env_path), env_path)
    rows = load_upload_rows(Path(args.workbook), limit=args.limit, start_row=args.start_row)
    output_dir = Path(args.output_dir)
    payload_dir = output_dir / "payloads"
    image_cache: dict[str, str] = {}
    log_rows: list[dict[str, Any]] = []

    for row in rows:
        now = dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds")
        log_row = {
            "timestamp_utc": now,
            "row": row["row"],
            "ean": row["ean"],
            "seller_sku": row["seller_sku"],
            "title": row["title"],
            "status": "",
            "product_id": "",
            "sku_id": "",
            "image_count": 0,
            "message": "",
        }
        try:
            if args.skip_existing:
                found = existing_sku(client, row["seller_sku"])
                if found:
                    log_row.update(status="skipped_existing_sku", product_id=found, message="Seller SKU already exists")
                    log_rows.append(log_row)
                    continue

            image_uris: list[str] = []
            for image_index, image_url in enumerate(row["image_urls"][: args.max_images], start=1):
                if image_url in image_cache:
                    image_uris.append(image_cache[image_url])
                    continue
                image_bytes = download_url(image_url)
                uploaded = client.upload_product_image(image_bytes, f"{row['ean']}_{image_index}.jpg")
                uri = clean(uploaded.get("uri"))
                if not uri:
                    raise TikTokApiError(f"Image upload returned no uri: {uploaded}")
                image_cache[image_url] = uri
                image_uris.append(uri)

            payload = build_payload(row, image_uris, args)
            payload_dir.mkdir(parents=True, exist_ok=True)
            (payload_dir / f"{row['ean']}.json").write_text(
                json.dumps(payload, ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
            log_row["image_count"] = len(image_uris)

            if args.dry_run:
                log_row.update(status="dry_run", message="Payload written, product not created")
            else:
                response = client.request(
                    "POST",
                    "/product/202309/products",
                    params={"shop_cipher": client.ensure_shop_cipher()},
                    body=payload,
                )
                data = response.get("data") or {}
                sku_rows = data.get("skus") or []
                log_row.update(
                    status="created",
                    product_id=clean(data.get("product_id")),
                    sku_id=clean((sku_rows[0] or {}).get("id")) if sku_rows else "",
                    message=compact_json(
                        {
                            "code": response.get("code"),
                            "message": response.get("message"),
                            "request_id": response.get("request_id"),
                            "warnings": data.get("warnings") or [],
                        }
                    ),
                )
        except Exception as exc:
            log_row.update(status="failed", message=str(exc)[:1500])
            log_rows.append(log_row)
            write_csv(output_dir / "api_product_upload_log.csv", log_rows)
            if args.stop_on_error:
                raise
            continue

        log_rows.append(log_row)
        write_csv(output_dir / "api_product_upload_log.csv", log_rows)

    print(f"Rows processed: {len(log_rows)}")
    print(f"Created: {sum(1 for row in log_rows if row['status'] == 'created')}")
    print(f"Skipped: {sum(1 for row in log_rows if row['status'].startswith('skipped'))}")
    print(f"Failed: {sum(1 for row in log_rows if row['status'] == 'failed')}")
    print(f"Log: {(output_dir / 'api_product_upload_log.csv').resolve()}")
    return 0


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Upload prepared Libri products to TikTok Shop through the API.")
    parser.add_argument("--env", default=".env")
    parser.add_argument("--workbook", default=str(DEFAULT_WORKBOOK))
    parser.add_argument("--output-dir", default=str(DEFAULT_OUTPUT_DIR))
    parser.add_argument("--limit", type=int, default=0)
    parser.add_argument("--start-row", type=int, default=7)
    parser.add_argument("--max-images", type=int, default=5)
    parser.add_argument("--save-mode", choices=["AS_DRAFT", "LISTING"], default="AS_DRAFT")
    parser.add_argument("--category-id", default=DEFAULT_CATEGORY_ID)
    parser.add_argument("--warehouse-id", default=DEFAULT_WAREHOUSE_ID)
    parser.add_argument("--responsible-person-id", default=DEFAULT_RESPONSIBLE_PERSON_ID)
    parser.add_argument("--currency", default=DEFAULT_CURRENCY)
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--skip-existing", action=argparse.BooleanOptionalAction, default=True)
    parser.add_argument("--stop-on-error", action=argparse.BooleanOptionalAction, default=True)
    return parser


def main(argv: list[str] | None = None) -> int:
    return process(build_parser().parse_args(argv))


if __name__ == "__main__":
    raise SystemExit(main())
